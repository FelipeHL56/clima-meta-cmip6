"""Genera el Excel del informe: proyecciones CMIP6 para el Meta (Colombia).

Lee los .nc de la carpeta de NetCDF (variable de entorno CMIP6_NC_DIR) (ver analisis_anomalias.py para el
método) y escribe Proyecciones_Meta_CMIP6.xlsx en .../resultados. Las anomalías por
modelo son entradas (azul); medias, rangos, desviaciones y tablas del resumen son fórmulas.
"""
import glob
import os
import re
from pathlib import Path

import numpy as np
import pandas as pd
import xarray as xr
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

PROY = Path(__file__).resolve().parents[1]
NC_DIR = Path(os.environ.get("CMIP6_NC_DIR", PROY / "datos" / "nc"))  # carpeta con los NetCDF (ruta sin tildes en Windows)
RES = PROY / "resultados" / "figuras"
OUT = PROY / "resultados" / "excel" / "Proyecciones_Meta_CMIP6.xlsx"

BASE = (1995, 2014)
PLAZOS = [("Corto", 2021, 2040), ("Mediano", 2041, 2060), ("Largo", 2081, 2100)]
ESC = [("ssp1_2_6", "SSP1-2.6 (bajas)", "2A78D6"), ("ssp2_4_5", "SSP2-4.5 (medias)", "EB6834"),
       ("ssp5_8_5", "SSP5-8.5 (altas)", "1BAF7A")]
VARS = {"near_surface_air_temperature": "tas", "precipitation": "pr"}
MODELOS = {  # id -> (nombre, institución, resolución nominal aprox.)
    "access_cm2": ("ACCESS-CM2", "CSIRO-ARCCSS (Australia)", "~1,9° x 1,25° (~200 km)"),
    "gfdl_esm4": ("GFDL-ESM4", "NOAA-GFDL (EE. UU.)", "~1° (~100 km)"),
    "mpi_esm1_2_lr": ("MPI-ESM1-2-LR", "MPI-M (Alemania)", "~1,9° (~200 km)"),
    "mri_esm2_0": ("MRI-ESM2-0", "MRI (Japón)", "~1,1° (~100 km)"),
    "noresm2_mm": ("NorESM2-MM", "NCC (Noruega)", "~1° (~100 km)"),
}
MIDS = list(MODELOS)
MES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

# ---------------- datos ----------------
mensual, celdas = {}, {}
for f in sorted(glob.glob(str(NC_DIR / "*.nc"))):
    m = re.match(r"(near_surface_air_temperature|precipitation)_(ssp\d_\d_\d|historical)_(.+)_(\d{4})-(\d{4})$",
                 os.path.basename(f)[:-3])
    var, exp, mod, _, _ = m.groups()
    v = VARS[var]
    ds = xr.open_dataset(f)
    da = ds[v].mean(["lat", "lon"])
    da = da * 86400 if v == "pr" else da - 273.15
    mensual[(v, exp, mod)] = pd.Series(da.values, index=pd.MultiIndex.from_arrays(
        [da.time.dt.year.values, da.time.dt.month.values]))
    celdas[mod] = (ds.lat.size, ds.lon.size)
anual = {k: s.groupby(level=0).mean() for k, s in mensual.items()}
base = {(v, m): anual[(v, "historical", m)].loc[BASE[0]:BASE[1]].mean() for v in VARS.values() for m in MIDS}


def anom(v, serie, m):
    return (serie / base[(v, m)] - 1) * 100 if v == "pr" else serie - base[(v, m)]


def clim(v, exp, m, a, b):
    s = mensual[(v, exp, m)]
    return s.loc[a:b].groupby(level=1).mean()


# ---------------- estilos ----------------
F = "Arial"
f_norm, f_b = Font(name=F, size=10), Font(name=F, size=10, bold=True)
f_in = Font(name=F, size=10, color="0000FF")
f_h = Font(name=F, size=10, bold=True, color="FFFFFF")
f_t = Font(name=F, size=14, bold=True, color="1F3864")
f_s = Font(name=F, size=10, italic=True, color="595959")
fill_h = PatternFill("solid", fgColor="1F3864")
fill_g = PatternFill("solid", fgColor="F2F2F2")
fill_k = {k: PatternFill("solid", fgColor=c) for k, _, c in ESC}
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_w = Alignment(horizontal="left", vertical="top", wrap_text=True)


def hdr(ws, row, col, vals, widths=None):
    for i, t in enumerate(vals):
        c = ws.cell(row, col + i, t)
        c.font, c.fill, c.alignment, c.border = f_h, fill_h, ctr, box


def put(ws, row, col, val, font=f_norm, fmt=None, fill=None, align=None):
    c = ws.cell(row, col, val)
    c.font, c.border = font, box
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    return c


def titulo(ws, t, sub=None):
    ws["A1"], ws["A1"].font = t, f_t
    if sub:
        ws["A2"], ws["A2"].font = sub, f_s
    ws.sheet_view.showGridLines = False


def pulir(ch, leyenda=True):
    """Evita superposiciones: títulos/leyenda fuera del área de trazado, etiquetas del eje X abajo."""
    ch.title.overlay = False
    ch.y_axis.title.overlay = False
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.x_axis.tickLblPos = "low"
    if leyenda and ch.legend is not None:
        ch.legend.position = "b"
        ch.legend.overlay = False
    return ch


wb = Workbook()

# ================= Temperatura / Precipitación =================
FILAS0 = 5
def hoja_var(nombre, v, unidad, fmt):
    ws = wb.create_sheet(nombre)
    titulo(ws, f"{nombre}: cambio respecto a 1995-2014 por modelo ({unidad})",
           "Azul = valor calculado a partir de los NetCDF (entrada). Negro = fórmula. Ventana de plazo menos línea base del mismo modelo.")
    cols = ["Escenario", "Plazo"] + [MODELOS[m][0] for m in MIDS] + ["Media ensamble", "Mínimo", "Máximo",
                                                                      "Desv. est.", "Modelos con aumento", "¿Señal > dispersión?"]
    hdr(ws, 4, 1, cols)
    r = FILAS0
    for k, nom, _ in ESC:
        for pn, a, b in PLAZOS:
            put(ws, r, 1, nom, f_b, fill=fill_k[k] if False else None)
            put(ws, r, 2, f"{pn} {a}-{b}")
            for j, m in enumerate(MIDS):
                val = float(anom(v, anual[(v, k, m)], m).loc[a:b].mean())
                put(ws, r, 3 + j, round(val, 4), f_in, fmt)
            rng = f"C{r}:G{r}"
            put(ws, r, 8, f"=AVERAGE({rng})", f_b, fmt)
            put(ws, r, 9, f"=MIN({rng})", fmt=fmt)
            put(ws, r, 10, f"=MAX({rng})", fmt=fmt)
            put(ws, r, 11, f"=STDEV({rng})", fmt=fmt)
            put(ws, r, 12, f'=COUNTIF({rng},">0")&" de "&COUNT({rng})', align=ctr)
            put(ws, r, 13, f'=IF(ABS(H{r})>K{r},"Sí","No: dentro de la dispersión")', align=ctr)
            r += 1
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 18
    for c in range(3, 14):
        ws.column_dimensions[CL(c)].width = 14
    ws.column_dimensions["L"].width = 20
    ws.column_dimensions["M"].width = 26
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "C5"
    r += 1
    ws.cell(r, 1, "Notas").font = f_b
    notas = [
        "Media ensamble = promedio simple de los 5 modelos. Desv. est. = dispersión entre modelos (n=5).",
        "'¿Señal > dispersión?' compara el valor absoluto de la media con la desviación entre modelos: si es 'No', el cambio no se distingue del ruido entre modelos.",
    ]
    if v == "pr":
        notas.append("Cambio relativo en % respecto a la precipitación media 1995-2014 del mismo modelo (los modelos no reproducen bien la precipitación absoluta local).")
    for t in notas:
        r += 1
        ws.cell(r, 1, t).font = f_s
    return ws


ws_t = hoja_var("Temperatura", "tas", "°C", "0.00")
ws_p = hoja_var("Precipitación", "pr", "%", "0.0")

# ================= Resumen =================
ws = wb.active
ws.title = "Resumen"
titulo(ws, "Proyecciones climáticas CMIP6 para el Meta (Colombia)",
       "Fuente: Copernicus Climate Data Store, dataset projections-cmip6. Ensamble de 5 modelos, cambio respecto a 1995-2014.")
ws["A4"], ws["A4"].font = "Hallazgos principales (se actualizan con las tablas)", f_b
T = lambda si, pi: f"Temperatura!H{FILAS0 + 3 * si + pi}"
P = lambda si, pi: f"'Precipitación'!H{FILAS0 + 3 * si + pi}"
def sg(ref):  # signo + ROUND: evita TEXT(), cuyo formato depende del idioma de Excel
    return f'IF({ref}>=0,"+","")&ROUND({ref},1)'


hall = [
    f'="Temperatura, largo plazo (2081-2100): "&{sg(T(0,2))}&" °C en SSP1-2.6, "&{sg(T(1,2))}&" °C en SSP2-4.5 y "&{sg(T(2,2))}&" °C en SSP5-8.5."',
    f'="Temperatura, corto plazo (2021-2040): entre "&{sg(T(0,0))}&" y "&{sg(T(2,0))}&" °C; los escenarios casi no se separan hasta ~2050."',
    f'="Precipitación, largo plazo: "&{sg(P(0,2))}&" % (SSP1-2.6), "&{sg(P(1,2))}&" % (SSP2-4.5) y "&{sg(P(2,2))}&" % (SSP5-8.5)."',
    '="Precipitación: "&IF(COUNTIF(\'Precipitación\'!M5:M13,"Sí")=0,"ninguna de las 9 combinaciones",COUNTIF(\'Precipitación\'!M5:M13,"Sí")&" de 9 combinaciones")&" tiene un cambio medio mayor que la dispersión entre modelos (ver columna final de la hoja Precipitación)."',
]
for i, t in enumerate(hall):
    c = ws.cell(5 + i, 1, t)
    c.font = f_norm

def tabla(ws, r0, titulo_t, fn, fmt):
    ws.cell(r0, 1, titulo_t).font = f_b
    hdr(ws, r0 + 1, 1, ["Plazo"] + [n for _, n, _ in ESC])
    for pi, (pn, a, b) in enumerate(PLAZOS):
        put(ws, r0 + 2 + pi, 1, f"{pn} {a}-{b}")
        for si in range(3):
            put(ws, r0 + 2 + pi, 2 + si, f"={fn(si, pi)}", f_norm, fmt, align=ctr)
    return r0 + 1

h1 = tabla(ws, 11, "Tabla 1. Cambio de temperatura media (°C) respecto a 1995-2014", T, "+0.0;-0.0;0.0")
h2 = tabla(ws, 18, "Tabla 2. Cambio de precipitación media (%) respecto a 1995-2014", P, "+0.0;-0.0;0.0")
ws.column_dimensions["A"].width = 24
for c in "BCD":
    ws.column_dimensions[c].width = 20


def barras(ws, h, tit, ytit, anchor):
    ch = BarChart()
    ch.type, ch.grouping = "col", "clustered"
    ch.title, ch.y_axis.title = tit, ytit
    for si, (_, nom, col) in enumerate(ESC):
        ch.add_data(Reference(ws, min_col=2 + si, min_row=h, max_row=h + 3), titles_from_data=True)
        ch.series[si].graphicalProperties.solidFill = col
        ch.series[si].graphicalProperties.line.solidFill = col
        ch.series[si].invertIfNegative = False
    ch.set_categories(Reference(ws, min_col=1, min_row=h + 1, max_row=h + 3))
    pulir(ch)
    ch.height, ch.width = 9.5, 17
    ws.add_chart(ch, anchor)


barras(ws, h1, "Cambio de temperatura por plazo y escenario", "°C", "F4")
barras(ws, h2, "Cambio de precipitación por plazo y escenario", "%", "F26")

# ================= Línea base =================
wl = wb.create_sheet("Línea base")
titulo(wl, "Línea base 1995-2014 (Historical) por modelo",
       "Promedio del recuadro del Meta. Azul = calculado de los NetCDF. Las anomalías se calculan contra la línea base del mismo modelo.")
hdr(wl, 4, 1, ["Modelo", "Temperatura (°C)", "Precipitación (mm/día)", "Precipitación (mm/año)"])
for i, m in enumerate(MIDS):
    r = 5 + i
    put(wl, r, 1, MODELOS[m][0])
    put(wl, r, 2, round(float(base[("tas", m)]), 3), f_in, "0.00")
    put(wl, r, 3, round(float(base[("pr", m)]), 3), f_in, "0.00")
    put(wl, r, 4, f"=C{r}*365", fmt="#,##0")
put(wl, 10, 1, "Media ensamble", f_b)
for c, fm in ((2, "0.00"), (3, "0.00"), (4, "#,##0")):
    put(wl, 10, c, f"=AVERAGE({CL(c)}5:{CL(c)}9)", f_b, fm)
put(wl, 11, 1, "Mín. – Máx. (rango)")
for c in (2, 3, 4):
    put(wl, 11, c, f'=ROUND(MIN({CL(c)}5:{CL(c)}9),1)&" – "&ROUND(MAX({CL(c)}5:{CL(c)}9),1)', align=ctr)
wl.cell(13, 1, "La dispersión de precipitación absoluta entre modelos es grande: por eso se reportan cambios relativos (%) y se recomienda corrección de sesgo contra ERA5 o estaciones.").font = f_s
for c, w in zip("ABCD", (24, 18, 22, 22)):
    wl.column_dimensions[c].width = w

# ================= Series anuales =================
def hoja_serie(nombre, v, ytxt, fmt):
    ws = wb.create_sheet(nombre)
    titulo(ws, f"{nombre}: anomalía anual suavizada (media móvil 10 años), respecto a 1995-2014",
           "Columnas azules = valor por modelo (entrada, calculado de los NetCDF). Media/Mín/Máx = fórmulas. Historical 1950-2014, SSP 2015-2100.")
    exps = [("historical", "Historical", "898781", 1950, 2014)] + [(k, n, c, 2015, 2100) for k, n, c in ESC]
    row_h = 27  # el gráfico ocupa las filas 4-25; la tabla va debajo
    hdr(ws, row_h, 1, ["Año"])
    col = 2
    bloques = {}
    for k, n, c, a, b in exps:
        bloques[k] = col
        for m in MIDS:
            hdr(ws, row_h, col, [f"{n} · {MODELOS[m][0]}"])
            col += 1
    stats = {}
    for k, n, c, a, b in exps:
        stats[k] = col
        hdr(ws, row_h, col, [f"{n} · Media", f"{n} · Mín", f"{n} · Máx"])
        col += 3
    anios = list(range(1950, 2101))
    for i, y in enumerate(anios):
        r = row_h + 1 + i
        put(ws, r, 1, y, f_b)
        for k, n, c, a, b in exps:
            if not (a <= y <= b):
                continue
            c0 = bloques[k]
            for j, m in enumerate(MIDS):
                s = anom(v, anual[(v, k, m)], m).rolling(10, center=True, min_periods=5).mean()
                put(ws, r, c0 + j, round(float(s.loc[y]), 4), f_in, fmt)
            rg = f"{CL(c0)}{r}:{CL(c0 + 4)}{r}"
            put(ws, r, stats[k], f"=AVERAGE({rg})", f_b, fmt)
            put(ws, r, stats[k] + 1, f"=MIN({rg})", fmt=fmt)
            put(ws, r, stats[k] + 2, f"=MAX({rg})", fmt=fmt)
    ws.column_dimensions["A"].width = 8
    for c in range(2, col):
        ws.column_dimensions[CL(c)].width = 16
    ws.row_dimensions[row_h].height = 42
    n = len(anios)
    ch = LineChart()
    ch.title, ch.y_axis.title = f"{nombre}: media del ensamble (media móvil 10 años)", ytxt
    for k, nom, c, a, b in exps:
        ch.add_data(Reference(ws, min_col=stats[k], min_row=row_h, max_row=row_h + n), titles_from_data=True)
        s = ch.series[-1]
        s.graphicalProperties.line.solidFill = c
        s.graphicalProperties.line.width = 22000
        s.smooth = False
    ch.set_categories(Reference(ws, min_col=1, min_row=row_h + 1, max_row=row_h + n))
    ch.x_axis.tickLblSkip = ch.x_axis.tickMarkSkip = 10
    pulir(ch)
    ch.height, ch.width = 10, 22
    ws.add_chart(ch, "A4")
    return ws


hoja_serie("Serie temperatura", "tas", "Δ °C", "0.00")
hoja_serie("Serie precipitación", "pr", "Δ %", "0.0")

# ================= Ciclo estacional =================
wc = wb.create_sheet("Ciclo estacional")
titulo(wc, "Ciclo estacional: línea base 1995-2014 y cambio mensual por escenario y plazo (media del ensamble)",
       "Valores calculados en Python (media de los 5 modelos, cada uno contra su propia línea base). El Meta tiene estación lluviosa ~abril-noviembre.")
hdr(wc, 4, 1, ["Mes", "Temp. base (°C)", "Precip. base (mm/día)"])
cb_t = pd.concat([clim("tas", "historical", m, *BASE) for m in MIDS], axis=1).mean(axis=1)
cb_p = pd.concat([clim("pr", "historical", m, *BASE) for m in MIDS], axis=1).mean(axis=1)
C0T, C0P = 5, 15
hdr(wc, 4, C0T, [f"Δ°C {n.split(' ')[0]} {pn}" for _, n, _ in ESC for pn, _, _ in PLAZOS])
hdr(wc, 4, C0P, [f"Δ% {n.split(' ')[0]} {pn}" for _, n, _ in ESC for pn, _, _ in PLAZOS])
wc.cell(3, C0T, "Cambio de temperatura (°C)").font = f_b
wc.cell(3, C0P, "Cambio de precipitación (%)").font = f_b
for mi in range(12):
    r = 5 + mi
    put(wc, r, 1, MES[mi], f_b)
    put(wc, r, 2, round(float(cb_t.loc[mi + 1]), 3), f_in, "0.00")
    put(wc, r, 3, round(float(cb_p.loc[mi + 1]), 3), f_in, "0.00")
    for si, (k, _, _) in enumerate(ESC):
        for pi, (pn, a, b) in enumerate(PLAZOS):
            dt = np.mean([clim("tas", k, m, a, b).loc[mi + 1] - clim("tas", "historical", m, *BASE).loc[mi + 1] for m in MIDS])
            dp = np.mean([(clim("pr", k, m, a, b).loc[mi + 1] / clim("pr", "historical", m, *BASE).loc[mi + 1] - 1) * 100 for m in MIDS])
            put(wc, r, C0T + 3 * si + pi, round(float(dt), 3), f_in, "0.00")
            put(wc, r, C0P + 3 * si + pi, round(float(dp), 2), f_in, "0.0")
put(wc, 17, 1, "Media anual", f_b)
for c in [2, 3] + list(range(C0T, C0T + 9)) + list(range(C0P, C0P + 9)):
    put(wc, 17, c, f"=AVERAGE({CL(c)}5:{CL(c)}16)", f_b, "0.00")
wc.cell(18, 1, "Nota: en los meses secos (dic-feb) la precipitación base es muy baja (~0,5-1 mm/día), por lo que los cambios porcentuales se inflan; en el informe conviene citar también el cambio absoluto.").font = f_s
wc.column_dimensions["A"].width = 14
for c in range(2, C0P + 9):
    wc.column_dimensions[CL(c)].width = 13
wc.row_dimensions[4].height = 42


def lineas_est(tit, ytit, c0, anchor, extra=None):
    ch = LineChart()
    ch.title, ch.y_axis.title = tit, ytit
    for si, (_, nom, col) in enumerate(ESC):
        ch.add_data(Reference(wc, min_col=c0 + 3 * si + 2, min_row=4, max_row=16), titles_from_data=True)
        s = ch.series[-1]
        s.graphicalProperties.line.solidFill = col
        s.graphicalProperties.line.width = 22000
        s.smooth = False
    ch.set_categories(Reference(wc, min_col=1, min_row=5, max_row=16))
    pulir(ch)
    ch.height, ch.width = 9.5, 17
    wc.add_chart(ch, anchor)


bc = BarChart()
bc.type = "col"
bc.title, bc.y_axis.title = "Precipitación media mensual 1995-2014 (ensamble)", "mm/día"
bc.add_data(Reference(wc, min_col=3, min_row=4, max_row=16), titles_from_data=True)
bc.series[0].graphicalProperties.solidFill = "2A78D6"
bc.set_categories(Reference(wc, min_col=1, min_row=5, max_row=16))
bc.legend = None
pulir(bc, leyenda=False)
bc.height, bc.width = 9.5, 17
wc.add_chart(bc, "A20")
lineas_est("Cambio mensual de temperatura, largo plazo (2081-2100)", "Δ °C", C0T, "J20")
lineas_est("Cambio mensual de precipitación, largo plazo (2081-2100)", "Δ %", C0P, "A41")

# ================= Comparación con observaciones =================
TAB = PROY / "resultados" / "tablas"
wo = wb.create_sheet("Comparación obs")
titulo(wo, "Modelos (Historical 1995-2014) frente a observaciones",
       "ERA5: mismo recuadro que los modelos. CHIRPS: promedio de la región. IDEAM: promedio de estaciones, temperatura media aproximada como (Tmax+Tmin)/2.")
cm = pd.read_csv(TAB / "comparacion_modelos_vs_observaciones.csv", encoding="utf-8-sig")
hdr(wo, 4, 1, ["Variable", "Referencia", "Modelo", "Media modelo", "Media observada", "Sesgo (modelo - obs.)", "Sesgo (%)",
               "RMSE ciclo anual", "Correlación ciclo anual", "Unidad"])
for i, r in cm.iterrows():
    row = 5 + i
    ens = r.modelo == "Ensamble"
    put(wo, row, 1, r.variable, f_b if ens else f_norm)
    put(wo, row, 2, r.referencia, f_b if ens else f_norm)
    put(wo, row, 3, r.modelo, f_b if ens else f_norm)
    put(wo, row, 4, float(r.media_modelo), f_in, "0.00")
    put(wo, row, 5, float(r.media_obs), f_in, "0.00")
    put(wo, row, 6, f"=D{row}-E{row}", f_b if ens else f_norm, "+0.00;-0.00;0.00")
    put(wo, row, 7, f"=F{row}/E{row}" if r.variable == "Precipitación" else "", f_b if ens else f_norm, "+0.0%;-0.0%;0.0%")
    put(wo, row, 8, float(r.RMSE_ciclo), f_in, "0.00")
    put(wo, row, 9, float(r.correl_ciclo), f_in, "0.00")
    put(wo, row, 10, r.unidad)
r0 = 5 + len(cm) + 2
wo.cell(r0, 1, "Climatología mensual 1995-2014 (media del ensamble de modelos y observaciones)").font = f_b
cl = pd.read_csv(TAB / "climatologia_modelos_y_observaciones_1995_2014.csv", index_col=0, encoding="utf-8-sig")
CLC = [("tas_Ensamble", "T modelos (°C)"), ("obs_tas_ERA5", "T ERA5 (°C)"), ("obs_tas_IDEAM", "T IDEAM aprox. (°C)"),
       ("pr_Ensamble", "P modelos (mm/mes)"), ("obs_pr_ERA5", "P ERA5 (mm/mes)"), ("obs_pr_CHIRPS", "P CHIRPS (mm/mes)"),
       ("obs_pr_IDEAM", "P IDEAM (mm/mes)")]
hdr(wo, r0 + 1, 1, ["Mes"] + [n for _, n in CLC])
for i, mes in enumerate(MES):
    put(wo, r0 + 2 + i, 1, mes, f_b)
    for j, (c, _) in enumerate(CLC):
        put(wo, r0 + 2 + i, 2 + j, float(cl.loc[mes, c]), f_in, "0.0")
put(wo, r0 + 14, 1, "Media/Total anual", f_b)
for j in range(len(CLC)):
    col = CL(2 + j)
    put(wo, r0 + 14, 2 + j, f"=AVERAGE({col}{r0 + 2}:{col}{r0 + 13})" if j < 3 else f"=SUM({col}{r0 + 2}:{col}{r0 + 13})", f_b, "0.0" if j < 3 else "#,##0")
wo.cell(r0 + 15, 1, "Temperatura: promedio anual (°C). Precipitación: total anual (mm/año). Azul = calculado de los datos (entrada).").font = f_s
for c, w in zip("ABCDEFGHIJ", (16, 14, 18, 14, 16, 18, 12, 16, 18, 10)):
    wo.column_dimensions[c].width = w
wo.row_dimensions[4].height = 32
wo.freeze_panes = "A5"
OBS_COLORS = {"ERA5": "2A78D6", "CHIRPS": "EB6834", "IDEAM": "1BAF7A"}


def lineas_obs(tit, ytit, cols, anchor):
    ch = LineChart()
    ch.title, ch.y_axis.title = tit, ytit
    for j in cols:
        ch.add_data(Reference(wo, min_col=2 + j, min_row=r0 + 1, max_row=r0 + 13), titles_from_data=True)
        s = ch.series[-1]
        nombre = CLC[j][1]
        col = "0B0B0B" if "modelos" in nombre else OBS_COLORS[nombre.split(" ")[1]]
        s.graphicalProperties.line.solidFill = col
        s.graphicalProperties.line.width = 28000 if "modelos" in nombre else 22000
        if "modelos" in nombre:
            s.graphicalProperties.line.dashStyle = "dash"
        s.smooth = False
    ch.set_categories(Reference(wo, min_col=1, min_row=r0 + 2, max_row=r0 + 13))
    pulir(ch)
    ch.height, ch.width = 9.5, 17
    wo.add_chart(ch, anchor)


lineas_obs("Ciclo estacional de temperatura 1995-2014", "°C", [0, 1, 2], "L4")
lineas_obs("Ciclo estacional de precipitación 1995-2014", "mm/mes", [3, 4, 5, 6], "L24")

# ================= Proyección ajustada (anual) =================
wp = wb.create_sheet("Proyección ajustada")
titulo(wp, "Proyecciones ajustadas con ERA5 (valores absolutos anuales por modelo)",
       "Temperatura: observado ERA5 + cambio del modelo. Precipitación: observado x factor trimestral móvil del modelo, renormalizado al cambio anual del modelo. Azul = entrada; negro = fórmula.")
an = pd.read_csv(TAB / "proyeccion_ajustada_anual.csv", encoding="utf-8-sig")
mod_names = [MODELOS[m][0] for m in MIDS]
hdr(wp, 4, 1, ["Variable", "Referencia", "Escenario", "Plazo", "Observado 1995-2014"] + mod_names +
    ["Media ensamble", "Mínimo", "Máximo", "Cambio vs observado", "Unidad anual"])
for i, r in an.iterrows():
    row = 5 + i
    tas = r.variable == "Temperatura"
    fm = "0.00" if tas else "#,##0"
    put(wp, row, 1, r.variable)
    put(wp, row, 2, r.referencia)
    put(wp, row, 3, r.escenario)
    put(wp, row, 4, r.plazo)
    put(wp, row, 5, float(r.observado_1995_2014), f_in, fm)
    for j, n in enumerate(mod_names):
        put(wp, row, 6 + j, float(r[n]), f_in, fm)
    put(wp, row, 11, f"=AVERAGE(F{row}:J{row})", f_b, fm)
    put(wp, row, 12, f"=MIN(F{row}:J{row})", fmt=fm)
    put(wp, row, 13, f"=MAX(F{row}:J{row})", fmt=fm)
    put(wp, row, 14, f"=K{row}-E{row}" if tas else f"=K{row}/E{row}-1", f_b, "+0.00;-0.00;0.00" if tas else "+0.0%;-0.0%;0.0%")
    put(wp, row, 15, r.unidad_anual)
last = 5 + len(an) - 1
for c, w in zip("ABCDEFGHIJKLMNO", (14, 11, 20, 18, 16, 13, 13, 15, 13, 13, 15, 11, 11, 16, 12)):
    wp.column_dimensions[c].width = w
wp.row_dimensions[4].height = 32
wp.freeze_panes = "E5"
rn = last + 3
wp.cell(rn, 1, "Matrices para los gráficos (referencia ERA5)").font = f_b


def matriz(r_ini, tit, base_row, fm):
    wp.cell(r_ini, 1, tit).font = f_b
    hdr(wp, r_ini + 1, 1, ["Plazo"] + [n for _, n, _ in ESC])
    put(wp, r_ini + 2, 1, "Observado 1995-2014")
    for si in range(3):
        put(wp, r_ini + 2, 2 + si, f"=$E${base_row}", fmt=fm, align=ctr)
    for pi, (pn, a, b) in enumerate(PLAZOS):
        put(wp, r_ini + 3 + pi, 1, f"{pn} {a}-{b}")
        for si in range(3):
            put(wp, r_ini + 3 + pi, 2 + si, f"=K{base_row + 3 * si + pi}", fmt=fm, align=ctr)
    return r_ini + 1


mt = matriz(rn + 1, "Temperatura media anual (°C), ERA5", 5, "0.00")
mp = matriz(rn + 8, "Precipitación anual (mm/año), ERA5", 5 + 9, "#,##0")


def barras_abs(h, tit, ytit, anchor, minimo=None):
    ch = BarChart()
    ch.type, ch.grouping = "col", "clustered"
    ch.title, ch.y_axis.title = tit, ytit
    for si, (_, nom, col) in enumerate(ESC):
        ch.add_data(Reference(wp, min_col=2 + si, min_row=h, max_row=h + 4), titles_from_data=True)
        ch.series[si].graphicalProperties.solidFill = col
        ch.series[si].graphicalProperties.line.solidFill = col
        ch.series[si].invertIfNegative = False
    ch.set_categories(Reference(wp, min_col=1, min_row=h + 1, max_row=h + 4))
    if minimo is not None:
        ch.y_axis.scaling.min = minimo
    pulir(ch)
    ch.height, ch.width = 9.5, 17
    wp.add_chart(ch, anchor)


barras_abs(mt, "Temperatura media anual proyectada (ajustada con ERA5)", "°C", "Q4", minimo=22)
barras_abs(mp, "Precipitación anual proyectada (ajustada con ERA5)", "mm/año", "Q24", minimo=0)

# ================= Proyección ajustada (mensual) =================
wq = wb.create_sheet("Proy. ajustada mensual")
titulo(wq, "Proyecciones ajustadas con ERA5 y CHIRPS: valores mensuales (media del ensamble y rango entre modelos)",
       "Valores calculados en Python (ver Metodología). Usa el filtro de la fila 4 para escoger variable, referencia, escenario o plazo.")
pm = pd.read_csv(TAB / "proyeccion_ajustada_mensual.csv", encoding="utf-8-sig")
hdr(wq, 4, 1, ["Variable", "Referencia", "Escenario", "Plazo", "Mes", "Observado 1995-2014", "Media ensamble", "Mínimo", "Máximo", "Unidad"])
for i, r in pm.iterrows():
    row = 5 + i
    for j, c in enumerate(["variable", "referencia", "escenario", "plazo", "mes"]):
        put(wq, row, 1 + j, r[c])
    for j, c in enumerate(["observado_1995_2014", "media_ensamble", "minimo", "maximo"]):
        put(wq, row, 6 + j, float(r[c]), f_in, "0.0")
    put(wq, row, 10, r.unidad)
wq.auto_filter.ref = f"A4:J{4 + len(pm)}"
wq.freeze_panes = "A5"
for c, w in zip("ABCDEFGHIJ", (14, 11, 20, 18, 8, 18, 15, 11, 11, 10)):
    wq.column_dimensions[c].width = w
wq.row_dimensions[4].height = 32

# ================= Sensibilidad de precipitación =================
wz = wb.create_sheet("Sensibilidad precip.")
titulo(wz, "Sensibilidad del método de ajuste de la precipitación (cambio anual, media del ensamble, base ERA5)",
       "El factor puramente mensual infla el cambio anual porque en los meses secos el modelo parte de una base casi nula.")
sn = pd.read_csv(TAB / "proyeccion_ajustada_sensibilidad_precipitacion.csv", encoding="utf-8-sig")
hdr(wz, 4, 1, ["Escenario", "Plazo", "a) Factor mensual", "b) Factor trimestral móvil", "c) Factor anual (= cambio del modelo)",
               "Adoptado: trimestral renormalizado", "Inflación del mensual (puntos)"])
for i, r in sn.iterrows():
    row = 5 + i
    put(wz, row, 1, r.escenario)
    put(wz, row, 2, r.plazo)
    put(wz, row, 3, float(r.iloc[2]) / 100, f_in, "+0.0%;-0.0%;0.0%")
    put(wz, row, 4, float(r.iloc[3]) / 100, f_in, "+0.0%;-0.0%;0.0%")
    put(wz, row, 5, float(r.iloc[4]) / 100, f_in, "+0.0%;-0.0%;0.0%")
    put(wz, row, 6, f"=E{row}", f_b, "+0.0%;-0.0%;0.0%")
    put(wz, row, 7, f"=(C{row}-F{row})*100", fmt="0.0")
wz.cell(5 + len(sn) + 1, 1, "Se adopta (f): el total anual cambia exactamente lo que cambia cada modelo; el factor trimestral solo redistribuye entre meses.").font = f_s
for c, w in zip("ABCDEFG", (20, 18, 16, 18, 22, 22, 20)):
    wz.column_dimensions[c].width = w
wz.row_dimensions[4].height = 44

# ================= ENSO (ONI) =================
PA = "'Proyección ajustada'"
we = wb.create_sheet("ENSO (ONI)")
titulo(we, "ENSO (ONI): efecto histórico sobre el Meta y comparación con el cambio climático",
       "Regresión mensual 1995-2025: anomalía ~ tendencia + ONI con rezago (errores Newey-West). Azul = entrada; negro = fórmula; fondo amarillo = supuesto editable.")
fill_y = PatternFill("solid", fgColor="FFFF00")
we.cell(4, 1, "1. Efecto histórico del ENSO (por cada 1 °C de ONI)").font = f_b
rg = pd.read_csv(TAB / "enso_efecto_regresion.csv", encoding="utf-8-sig")
hdr(we, 5, 1, ["Variable", "Rezago (meses)", "Correlación (sin tendencia)", "Efecto por 1 °C de ONI", "Error (Newey-West)", "t",
               "Tendencia por década", "Unidad", "¿|t| > 2?"])
for i, r in rg.iterrows():
    row = 6 + i
    put(we, row, 1, r.variable)
    put(we, row, 2, int(r.rezago_meses), f_in, "0", align=ctr)
    put(we, row, 3, float(r.corr_detrend), f_in, "0.00")
    put(we, row, 4, float(r.efecto_por_1C_ONI), f_in, "+0.00;-0.00;0.00")
    put(we, row, 5, float(r.error_NW), f_in, "0.00")
    put(we, row, 6, float(r.t), f_in, "0.0")
    put(we, row, 7, float(r.tendencia_por_decada), f_in, "+0.00;-0.00;0.00")
    put(we, row, 8, r.unidad)
    put(we, row, 9, f'=IF(ABS(F{row})>2,"Sí","No")', align=ctr)
we.cell(12, 1, "2. Supuestos de un año con ENSO sostenido (editables)").font = f_b
put(we, 13, 1, "ONI sostenido, El Niño fuerte (°C)")
put(we, 13, 2, 1.5, f_in, "0.0", fill=fill_y)
put(we, 14, 1, "ONI sostenido, La Niña (°C)")
put(we, 14, 2, -1.0, f_in, "0.0", fill=fill_y)
put(we, 15, 1, "Efecto en T anual, El Niño fuerte (°C)")
put(we, 15, 2, "=D6*B13", f_b, "+0.00;-0.00;0.00")
put(we, 16, 1, "Efecto en T anual, La Niña (°C)")
put(we, 16, 2, "=D6*B14", f_b, "+0.00;-0.00;0.00")
put(we, 17, 1, "Incertidumbre del efecto El Niño (± °C, 1,96 x error)")
put(we, 17, 2, "=1.96*E6*B13", fmt="0.00")
we.cell(18, 1, "Referencia: 2015 tuvo ONI medio +1,55 (El Niño fuerte). El efecto usa la regresión de T ERA5 (fila 6) y supone que la relación y la amplitud del ENSO no cambian.").font = f_s
we.cell(20, 1, "3. Temperatura anual proyectada con efecto ENSO (ERA5, media del ensamble)").font = f_b
hdr(we, 21, 1, ["Escenario", "Plazo", "Cambio climático vs 1995-2014 (°C)", "T proyectada, año neutro (°C)",
                "T en año El Niño fuerte (°C)", "T en año La Niña (°C)", "Etiqueta"])
for k in range(9):
    row, src = 22 + k, 5 + k
    put(we, row, 1, f"={PA}!C{src}")
    put(we, row, 2, f"={PA}!D{src}")
    put(we, row, 3, f"={PA}!N{src}", fmt="+0.00;-0.00;0.00")
    put(we, row, 4, f"={PA}!K{src}", fmt="0.00")
    put(we, row, 5, f"=D{row}+$B$15", f_b, "0.00")
    put(we, row, 6, f"=D{row}+$B$16", f_b, "0.00")
    put(we, row, 7, f'=LEFT(A{row},FIND(" ",A{row})-1)&" "&LEFT(B{row},FIND(" ",B{row})-1)')
we.cell(32, 1, "4. Años con simulación WRF en el contexto ENSO").font = f_b
hdr(we, 33, 1, ["Año", "ONI medio anual", "ONI máximo", "ONI mínimo", "Fase (|ONI medio| ≥ 0,5)"])
ay = pd.read_csv(TAB / "enso_anios_wrf_oni.csv", encoding="utf-8-sig")
for i, r in ay.iterrows():
    row = 34 + i
    put(we, row, 1, int(r.anio), f_b, "0", align=ctr)
    put(we, row, 2, float(r.ONI_medio_anual), f_in, "+0.00;-0.00;0.00")
    put(we, row, 3, float(r.ONI_max), f_in, "+0.00;-0.00;0.00")
    put(we, row, 4, float(r.ONI_min), f_in, "+0.00;-0.00;0.00")
    put(we, row, 5, f'=IF(B{row}>=0.5,"El Niño",IF(B{row}<=-0.5,"La Niña","Neutro"))', align=ctr)
we.cell(42, 1, "5. Variabilidad interanual 1995-2014 (sin tendencia): modelos frente a observaciones").font = f_b
hdr(we, 43, 1, ["Fuente", "Desv. est. T anual (°C)", "Desv. est. P anual (%)", "T: veces la de ERA5", "P: veces la de ERA5"])
vr = pd.read_csv(TAB / "variabilidad_interanual_modelos_vs_obs.csv", encoding="utf-8-sig")
nombres_v = {"ERA5": "ERA5", "CHIRPS": "CHIRPS"}
nombres_v.update({m: MODELOS[m][0] for m in MIDS})
for i, r in vr.iterrows():
    row = 44 + i
    put(we, row, 1, nombres_v.get(r.fuente, r.fuente), f_b if i < 2 else f_norm)
    put(we, row, 2, float(r.std_T_C) if pd.notna(r.std_T_C) else "", f_in, "0.00")
    put(we, row, 3, float(r["std_P_%"]), f_in, "0.0")
    put(we, row, 4, f"=B{row}/$B$44" if pd.notna(r.std_T_C) else "", fmt="0.0")
    put(we, row, 5, f"=C{row}/$C$44", fmt="0.0")
we.cell(44 + len(vr) + 1, 1, "Los modelos muestran más variabilidad interanual que ERA5 (muestra de 20 años); su ENSO simulado no coincide en el tiempo con el real.").font = f_s
for c, w in zip("ABCDEFGHI", (44, 20, 24, 24, 24, 20, 20, 20, 12)):
    we.column_dimensions[c].width = w
we.row_dimensions[5].height = 44
we.row_dimensions[21].height = 44
ch = BarChart()
ch.type, ch.grouping = "col", "clustered"
ch.title, ch.y_axis.title = "Temperatura anual proyectada en año neutro, El Niño fuerte y La Niña", "°C"
for j, col in ((4, "898781"), (5, "E34948"), (6, "2A78D6")):
    ch.add_data(Reference(we, min_col=j, min_row=21, max_row=30), titles_from_data=True)
    ch.series[-1].graphicalProperties.solidFill = col
    ch.series[-1].graphicalProperties.line.solidFill = col
ch.set_categories(Reference(we, min_col=7, min_row=22, max_row=30))
ch.y_axis.scaling.min = 22
pulir(ch)
ch.height, ch.width = 10, 20
we.add_chart(ch, "K4")

# ================= WRF =================
ww = wb.create_sheet("WRF")
titulo(ww, "WRF (4 km): validación frente a ERA5 e IDEAM y zonas cordillera/llanos",
       "WRF NCAR RDA d616000, temperatura del nivel más bajo, 7 años (2000, 2011, 2013-2016, 2018). Recuadro lat 2-5, lon -74,5 a -71. Cordillera = celdas con media 2000-2018 < 22 °C.")
wa = pd.read_csv(TAB / "wrf_anual_vs_era5_oni.csv", encoding="utf-8-sig")
ww.cell(4, 1, "1. Temperatura media anual por año simulado (°C)").font = f_b
hdr(ww, 5, 1, ["Año", "Fase ENSO", "ONI medio", "WRF recuadro", "ERA5", "Sesgo WRF - ERA5", "WRF llanos", "IDEAM aprox.",
               "Sesgo llanos - IDEAM", "WRF cordillera"])
for i, r in wa.iterrows():
    row = 6 + i
    put(ww, row, 1, int(r.anio), f_b, "0", align=ctr)
    put(ww, row, 2, f'=IF(C{row}>=0.5,"El Niño",IF(C{row}<=-0.5,"La Niña","Neutro"))', align=ctr)
    put(ww, row, 3, float(r.ONI), f_in, "+0.00;-0.00;0.00")
    put(ww, row, 4, float(r.recuadro), f_in, "0.00")
    put(ww, row, 5, float(r.ERA5), f_in, "0.00")
    put(ww, row, 6, f"=D{row}-E{row}", fmt="+0.00;-0.00;0.00")
    put(ww, row, 7, float(r.llanos), f_in, "0.00")
    put(ww, row, 8, float(r.IDEAM_aprox), f_in, "0.00")
    put(ww, row, 9, f"=G{row}-H{row}", fmt="+0.00;-0.00;0.00")
    put(ww, row, 10, float(r.cordillera), f_in, "0.00")
put(ww, 13, 1, "Media", f_b)
for c in range(3, 11):
    put(ww, 13, c, f"=AVERAGE({CL(c)}6:{CL(c)}12)", f_b, "+0.00;-0.00;0.00" if c in (6, 9) else "0.00")
put(ww, 14, 1, "Correl. WRF-ERA5 (7 años)")
put(ww, 14, 4, "=CORREL(D6:D12,E6:E12)", fmt="0.00")
ww.cell(15, 1, "La correlación de 0,94 es indicativa (n=7) y en parte refleja la tendencia de calentamiento compartida.").font = f_s
ww.cell(17, 1, "2. Métricas mensuales (84 meses)").font = f_b
hdr(ww, 18, 1, ["Comparación", "Sesgo (°C)", "RMSE (°C)", "Correlación mensual", "Correl. de anomalías mensuales"])
mt_ = pd.read_csv(TAB / "wrf_validacion_metricas.csv", encoding="utf-8-sig")
for i, r in mt_.iterrows():
    row = 19 + i
    put(ww, row, 1, r.comparacion)
    put(ww, row, 2, float(r.sesgo_C), f_in, "+0.00;-0.00;0.00")
    put(ww, row, 3, float(r.RMSE_C), f_in, "0.00")
    put(ww, row, 4, float(r.correlacion_mensual), f_in, "0.00")
    put(ww, row, 5, float(r.correl_anomalia_mensual), f_in, "0.00")
ww.cell(22, 1, "3. Zonas del recuadro (WRF)").font = f_b
hdr(ww, 23, 1, ["Zona", "Celdas de 4 km", "% del recuadro", "T media anual WRF (°C)"])
zc = pd.read_csv(TAB / "wrf_zonas_celdas.csv", encoding="utf-8-sig").set_index("zona")
for i, (zn, nom, colT) in enumerate((("cordillera", "Cordillera (< 22 °C)", "J"), ("llanos", "Llanos (≥ 22 °C)", "G"), ("recuadro", "Recuadro completo", "D"))):
    row = 24 + i
    put(ww, row, 1, nom)
    put(ww, row, 2, int(zc.loc[zn, "celdas"]), f_in, "#,##0")
    put(ww, row, 3, f"=B{row}/$B$26", fmt="0.0%")
    put(ww, row, 4, f"={colT}13", fmt="0.00")
ww.cell(28, 1, "4. Serie mensual (°C)").font = f_b
hdr(ww, 29, 1, ["Periodo", "WRF recuadro", "ERA5", "Sesgo", "WRF llanos", "IDEAM aprox.", "Sesgo", "WRF cordillera"])
vm = pd.read_csv(TAB / "wrf_validacion_mensual.csv", index_col=0, encoding="utf-8-sig")
for i, (per, r) in enumerate(vm.iterrows()):
    row = 30 + i
    put(ww, row, 1, str(per), f_b)
    put(ww, row, 2, float(r.WRF_recuadro), f_in, "0.00")
    put(ww, row, 3, float(r.ERA5), f_in, "0.00")
    put(ww, row, 4, f"=B{row}-C{row}", fmt="+0.00;-0.00;0.00")
    put(ww, row, 5, float(r.WRF_llanos), f_in, "0.00")
    put(ww, row, 6, float(r.IDEAM_aprox), f_in, "0.00")
    put(ww, row, 7, f"=E{row}-F{row}", fmt="+0.00;-0.00;0.00")
    put(ww, row, 8, float(r.WRF_cordillera), f_in, "0.00")
last_w = 30 + len(vm) - 1
for c, w in zip("ABCDEFGHIJ", (26, 14, 14, 16, 14, 16, 14, 16, 18, 16)):
    ww.column_dimensions[c].width = w
ww.row_dimensions[5].height = 32
ww.row_dimensions[18].height = 32
ww.row_dimensions[29].height = 32
b2 = BarChart()
b2.type, b2.grouping = "col", "clustered"
b2.title, b2.y_axis.title = "Temperatura media anual del recuadro: WRF y ERA5", "°C"
for j, col in ((4, "0B0B0B"), (5, "2A78D6")):
    b2.add_data(Reference(ww, min_col=j, min_row=5, max_row=12), titles_from_data=True)
    b2.series[-1].graphicalProperties.solidFill = col
b2.set_categories(Reference(ww, min_col=1, min_row=6, max_row=12))
b2.y_axis.scaling.min = 22
pulir(b2)
b2.height, b2.width = 9, 16
ww.add_chart(b2, "L4")
l2 = LineChart()
l2.title, l2.y_axis.title = "Serie mensual del recuadro: WRF frente a ERA5 (84 meses)", "°C"
for j, col in ((2, "0B0B0B"), (3, "2A78D6")):
    l2.add_data(Reference(ww, min_col=j, min_row=29, max_row=last_w), titles_from_data=True)
    l2.series[-1].graphicalProperties.line.solidFill = col
    l2.series[-1].graphicalProperties.line.width = 18000
    l2.series[-1].smooth = False
l2.set_categories(Reference(ww, min_col=1, min_row=30, max_row=last_w))
l2.x_axis.tickLblSkip = l2.x_axis.tickMarkSkip = 12
l2.y_axis.scaling.min, l2.y_axis.scaling.max = 20, 29
pulir(l2)
l2.height, l2.width = 9, 22
ww.add_chart(l2, "L23")

# ================= Zonas cordillera / llanos =================
wzo = wb.create_sheet("Zonas")
titulo(wzo, "Temperatura proyectada por zona: llanos, recuadro y cordillera",
       "Proyección ajustada con ERA5 para el recuadro + diferencia espacial del WRF (media 2000-2018). Supuesto: mismo calentamiento en toda la zona.")
wzo.cell(4, 1, "1. Diferencia de cada zona respecto al recuadro (°C, media mensual WRF)").font = f_b
hdr(wzo, 5, 1, ["Mes", "Cordillera - recuadro", "Llanos - recuadro"])
of = pd.read_csv(TAB / "wrf_offsets_zona_vs_recuadro.csv", index_col=0, encoding="utf-8-sig")
for i, mes in enumerate(MES):
    put(wzo, 6 + i, 1, mes, f_b)
    put(wzo, 6 + i, 2, float(of.loc[mes, "cordillera"]), f_in, "0.00")
    put(wzo, 6 + i, 3, float(of.loc[mes, "llanos"]), f_in, "0.00")
put(wzo, 18, 1, "Media anual", f_b)
put(wzo, 18, 2, "=AVERAGE(B6:B17)", f_b, "0.00")
put(wzo, 18, 3, "=AVERAGE(C6:C17)", f_b, "0.00")
wzo.cell(21, 1, "2. Temperatura media anual por zona (°C)").font = f_b
hdr(wzo, 22, 1, ["Escenario", "Plazo", "Etiqueta", "Recuadro", "Cordillera", "Llanos",
                 "Mín. recuadro", "Mín. cordillera", "Mín. llanos", "Máx. recuadro", "Máx. cordillera", "Máx. llanos", "Cambio vs 1995-2014"])
put(wzo, 23, 1, "Observado ERA5 + diferencia WRF", f_b)
put(wzo, 23, 2, "1995-2014")
put(wzo, 23, 3, "Observado")
put(wzo, 23, 4, f"={PA}!E5", f_b, "0.00")
put(wzo, 23, 5, "=D23+$B$18", f_b, "0.00")
put(wzo, 23, 6, "=D23+$C$18", f_b, "0.00")
for c in range(7, 13):
    put(wzo, 23, c, "")
put(wzo, 23, 13, "")
for k in range(9):
    row, src = 24 + k, 5 + k
    put(wzo, row, 1, f"={PA}!C{src}")
    put(wzo, row, 2, f"={PA}!D{src}")
    put(wzo, row, 3, f'=LEFT(A{row},FIND(" ",A{row})-1)&" "&LEFT(B{row},FIND(" ",B{row})-1)')
    put(wzo, row, 4, f"={PA}!K{src}", f_b, "0.00")
    put(wzo, row, 5, f"=D{row}+$B$18", f_b, "0.00")
    put(wzo, row, 6, f"=D{row}+$C$18", f_b, "0.00")
    put(wzo, row, 7, f"={PA}!L{src}", fmt="0.00")
    put(wzo, row, 8, f"=G{row}+$B$18", fmt="0.00")
    put(wzo, row, 9, f"=G{row}+$C$18", fmt="0.00")
    put(wzo, row, 10, f"={PA}!M{src}", fmt="0.00")
    put(wzo, row, 11, f"=J{row}+$B$18", fmt="0.00")
    put(wzo, row, 12, f"=J{row}+$C$18", fmt="0.00")
    put(wzo, row, 13, f"=D{row}-$D$23", fmt="+0.00;-0.00;0.00")
wzo.cell(34, 1, "El cambio es igual en las tres zonas por construcción. La cordillera podría calentarse algo más (calentamiento dependiente de la altura): no se evalúa aquí.").font = f_s
wzo.cell(35, 1, "Referencia: con SSP5-8.5 a largo plazo, la zona de cordillera (media ~12 °C) pasaría a ~17 °C y los llanos (~25 °C) a ~30 °C.").font = f_s
for c, w in zip("ABCDEFGHIJKLM", (30, 18, 16, 12, 12, 12, 13, 14, 12, 13, 14, 12, 16)):
    wzo.column_dimensions[c].width = w
wzo.row_dimensions[22].height = 44
b3 = BarChart()
b3.type, b3.grouping = "col", "clustered"
b3.title, b3.y_axis.title = "Temperatura media anual proyectada por zona (media del ensamble)", "°C"
for j, col in ((5, "2A78D6"), (4, "898781"), (6, "EB6834")):
    b3.add_data(Reference(wzo, min_col=j, min_row=22, max_row=32), titles_from_data=True)
    b3.series[-1].graphicalProperties.solidFill = col
    b3.series[-1].graphicalProperties.line.solidFill = col
b3.set_categories(Reference(wzo, min_col=3, min_row=23, max_row=32))
pulir(b3)
b3.height, b3.width = 10, 22
wzo.add_chart(b3, "O4")

# ================= Modelos =================
wm = wb.create_sheet("Modelos")
titulo(wm, "Modelos del ensamble", "Un solo miembro (r1i1p1f1) por modelo. Celdas = malla dentro del recuadro N 5, S 2, O -74,5, E -71.")
hdr(wm, 4, 1, ["Modelo", "Institución", "Resolución nominal", "Celdas en el recuadro (lat x lon)", "N.º de celdas",
               "Temperatura (3 SSP + Historical)", "Precipitación (3 SSP + Historical)"])
for i, m in enumerate(MIDS):
    r = 5 + i
    put(wm, r, 1, MODELOS[m][0], f_b)
    put(wm, r, 2, MODELOS[m][1])
    put(wm, r, 3, MODELOS[m][2])
    put(wm, r, 4, f"{celdas[m][0]} x {celdas[m][1]}", f_in, align=ctr)
    put(wm, r, 5, f"={celdas[m][0]}*{celdas[m][1]}", fmt="0", align=ctr)
    put(wm, r, 6, "Completo", align=ctr)
    put(wm, r, 7, "Completo", align=ctr)
wm.cell(11, 1, "Descartado: EC-Earth3 (el CDS devolvió RoocsValueError en todas las solicitudes, incluso sin recorte espacial).").font = f_s
wm.cell(12, 1, "Resolución nominal aproximada, referencial; la malla efectiva es la de la columna 'Celdas'.").font = f_s
for c, w in zip("ABCDEFG", (18, 28, 26, 24, 14, 26, 26)):
    wm.column_dimensions[c].width = w
wm.row_dimensions[4].height = 42

# ================= Metodología =================
wd = wb.create_sheet("Metodología")
titulo(wd, "Metodología y supuestos")
filas = [
    ("Fuente", "Copernicus Climate Data Store (C3S), dataset 'CMIP6 climate projections' (projections-cmip6). Descarga con la API cdsapi."),
    ("Zona de estudio", "Región del Meta. Recuadro: Norte 5,0°, Sur 2,0°, Oeste -74,5°, Este -71,0°. Solo entran celdas cuyo centro cae dentro del recuadro (2 a 9 según el modelo)."),
    ("Variables", "Temperatura del aire cerca de la superficie (tas) y precipitación (pr). Resolución temporal mensual."),
    ("Experimentos", "Bajas emisiones: SSP1-2.6. Medias: SSP2-4.5. Altas: SSP5-8.5. Referencia: Historical (1850-2014). Series SSP 2015-2100."),
    ("Modelos", "NorESM2-MM, MRI-ESM2-0, MPI-ESM1-2-LR, GFDL-ESM4 y ACCESS-CM2 (miembro r1i1p1f1). Criterios: disponibilidad de todos los experimentos y variables, diversidad institucional, evitar solo modelos 'muy calientes', resolución."),
    ("Plazos", "Corto 2021-2040, Mediano 2041-2060, Largo 2081-2100. Ventanas de 20 años (estándar IPCC) para separar la señal de la variabilidad natural (El Niño y La Niña)."),
    ("Línea base", "Historical 1995-2014, calculada por modelo."),
    ("Cálculo", "Promedio simple de las celdas del recuadro, media anual por modelo, media de la ventana menos la línea base del mismo modelo. Temperatura: diferencia en °C. Precipitación: cambio relativo (%)."),
    ("Ensamble", "Media simple de los 5 modelos (sin ponderar). La dispersión se reporta como mínimo, máximo y desviación estándar entre modelos."),
    ("Unidades", "tas: K convertida a °C. pr: kg m-2 s-1 convertida a mm/día (x 86 400)."),
    ("Series suavizadas", "Hojas 'Serie ...': media móvil centrada de 10 años (mínimo 5 años en los extremos) sobre la anomalía anual."),
    ("Limitaciones", "Mallas gruesas (~100-200 km): sirven para la tendencia regional, no para municipios. Un solo miembro por modelo (variabilidad interna). No hay corrección de sesgo; la precipitación absoluta de los modelos difiere de la observada. Se recomienda comparar con ERA5 o estaciones."),
    ("Observaciones", "ERA5 (T y P mensuales, mismo recuadro lat 2-5, lon -74,5 a -71), CHIRPS (P mensual, promedio de la región) e IDEAM (promedio de estaciones: 14 de lluvia y 7-10 de temperatura; solo Tmax y Tmin, así que la temperatura media es (Tmax+Tmin)/2, aproximada). Periodo de comparación 1995-2014. Datos en la carpeta del usuario 'Datos modelación'."),
    ("Ajuste con observaciones", "Factor de cambio (método delta). Temperatura: observado ERA5 de cada mes + cambio mensual del modelo. Precipitación: observado x factor trimestral móvil del modelo (escenario/histórico), renormalizado para que el total anual cambie lo mismo que en el modelo. El factor puramente mensual se descartó: infla el cambio anual 2-6 puntos por los meses secos (hoja 'Sensibilidad precip.'). Referencia principal ERA5; CHIRPS como contraste en lluvia; MPI-ESM1-2-LR se mantiene (es el de mayor sesgo en lluvia, -42 % a -50 %)."),
    ("ENSO (ONI)", "ONI mensual (NOAA, Niño 3.4, media móvil de 3 meses) de 1950 a jul-2026. Fases: El Niño ONI ≥ +0,5; La Niña ≤ -0,5. Efecto sobre el Meta: regresión mensual 1995-2025 de la anomalía (sin la climatología del mes) sobre tendencia lineal + ONI con rezago de máxima correlación (0-6 meses), con errores estándar Newey-West (12 rezagos) por la autocorrelación mensual. Para la proyección se supone que el efecto por °C de ONI y la amplitud del ENSO no cambian en el futuro (CMIP6 no permite fijar esto con confianza)."),
    ("WRF", "Simulación WRF de 4 km (NCAR RDA d616000), temperatura del nivel más bajo, 7 años: 2000 y 2011 (La Niña), 2013, 2014, 2016 y 2018 (neutros por ONI medio anual) y 2015 (El Niño fuerte). Medias mensuales calculadas del NetCDF horario sobre el recuadro lat 2-5, lon -74,5 a -71. Cordillera = celdas con media 2000-2018 < 22 °C (13 % del recuadro); llanos = el resto."),
    ("Zonas", "Proyección por zona = proyección ajustada con ERA5 del recuadro + diferencia mensual de cada zona respecto al recuadro en el WRF (media de los 7 años). Supone un calentamiento uniforme: el WRF aporta solo el contraste espacial, no un downscaling dinámico del escenario futuro."),
    ("Convención de colores", "Azul = valor de entrada calculado a partir de los datos. Negro = fórmula de Excel. Fondo amarillo = supuesto editable."),
]
hdr(wd, 3, 1, ["Tema", "Descripción"])
for i, (a, b) in enumerate(filas):
    put(wd, 4 + i, 1, a, f_b, align=left_w)
    put(wd, 4 + i, 2, b, align=left_w)
    wd.row_dimensions[4 + i].height = 15 * (1 + len(b) // 110)
wd.column_dimensions["A"].width = 22
wd.column_dimensions["B"].width = 120

# ================= Figuras =================
wf = wb.create_sheet("Figuras")
titulo(wf, "Figuras (PNG) para el informe", "Los mismos gráficos están en la carpeta resultados.")
r = 4
for fn, w in (("anomalia_tas_por_plazo.png", 800), ("anomalia_pr_por_plazo.png", 800), ("series_temporales_ensamble.png", 1000),
              ("comp_ciclo_estacional_modelos_vs_obs.png", 1000), ("comp_sesgo_modelos.png", 1000),
              ("comp_observaciones_entre_si.png", 1000), ("proy_resumen_anual_absoluto.png", 1000),
              ("proy_ciclo_estacional_pr.png", 1000), ("proy_ciclo_estacional_tas.png", 1000),
              ("enso_oni_serie_y_anios_wrf.png", 1000), ("enso_efecto_temperatura_lluvia.png", 1000),
              ("enso_vs_cambio_climatico.png", 900), ("wrf_validacion_mensual.png", 1000),
              ("wrf_mapa_temperatura_y_malla_cmip6.png", 800), ("wrf_anual_vs_era5_enso.png", 1000),
              ("zonas_proyeccion_temperatura.png", 1000)):
    p = RES / fn
    if p.exists():
        img = Image(str(p))
        img.height, img.width = img.height * w / img.width, w
        wf.add_image(img, f"A{r}")
        r += int(img.height / 20) + 3

orden = ["Resumen", "Temperatura", "Precipitación", "Serie temperatura", "Serie precipitación", "Ciclo estacional",
         "Comparación obs", "Proyección ajustada", "Proy. ajustada mensual", "Sensibilidad precip.",
         "ENSO (ONI)", "WRF", "Zonas", "Línea base", "Modelos", "Metodología", "Figuras"]
wb._sheets = [wb[n] for n in orden]
wb.save(OUT)
print("Guardado", OUT)

